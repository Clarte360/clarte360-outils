from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from collections import Counter
from pathlib import Path

import openpyxl

from clarte360_pip.pip_data.loader import load_pip_bank

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "docs/sources/CLARTE360_PIP_RIASEC_TABLEUR_MAITRE_V0_6_ACTIF.xlsx"
RUNTIME = ROOT / "resources/runtime/pip_bank_PIP-BANK-0.5.json"


def test_master_has_only_three_visible_working_sheets():
    wb = openpyxl.load_workbook(MASTER, read_only=False, data_only=False)
    visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    assert visible == ["01_QUESTIONS_72", "02_REPONSES_CONSIGNE", "03_REVUE_HISTORIQUE"]
    assert wb.active.title == "01_QUESTIONS_72"


def test_master_72_contract_and_interest_wording():
    wb = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
    ws = wb["01_QUESTIONS_72"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[idx["ordre"]] is not None]
    assert len(rows) == 72
    assert Counter(r[idx["dimension"]] for r in rows) == Counter({d: 12 for d in "RIASEC"})
    facets = Counter(r[idx["facette"]] for r in rows)
    assert len(facets) == 30 and min(facets.values()) >= 2
    assert {r[idx["item_version"]] for r in rows} == {"0.6"}
    assert {r[idx["bank_version"]] for r in rows} == {"PIP-BANK-0.5"}
    assert {r[idx["statut"]] for r in rows} == {"ACTIF"}
    assert all(len(str(r[idx["question_finale"]]).strip()) >= 20 for r in rows)


def test_response_protocol_is_interest_not_competence():
    wb = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
    ws = wb["02_REPONSES_CONSIGNE"]
    instruction = ws["B4"].value
    assert "aimeriez" in instruction.lower()
    assert "savez déjà faire" in instruction.lower()
    labels = [ws.cell(r, 2).value for r in range(8, 13)]
    assert labels == [
        "Je n’aimerais pas du tout faire cela",
        "J’aimerais peu faire cela",
        "Je suis partagé(e) / sans préférence marquée",
        "J’aimerais assez faire cela",
        "J’aimerais beaucoup faire cela",
    ]


def test_runtime_is_generated_from_master_and_in_sync():
    result = subprocess.run([sys.executable, str(ROOT / "scripts/build_pip_runtime.py"), "--check"], cwd=ROOT, text=True, capture_output=True)
    assert result.returncode == 0, result.stdout + result.stderr
    bank = load_pip_bank()
    assert bank["source_file"] == MASTER.name
    assert bank["source_sha256"] == hashlib.sha256(MASTER.read_bytes()).hexdigest()
    assert bank["bank_version"] == "PIP-BANK-0.5"
    assert len(bank["items"]) == 72


def test_runtime_contains_no_separate_concrete_example_field():
    bank = json.loads(RUNTIME.read_text(encoding="utf-8"))
    assert all("exemple_concret" not in item for item in bank["items"])
