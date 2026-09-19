from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/sources/REFERENTIEL_INTERPRETATION_PIP_RIASEC_CLARTE360_V1_0.xlsx'

NAVY='183B56'; TEAL='008080'; LIGHT='EAF4F4'; PALE='F4F7F8'; WHITE='FFFFFF'; GREY='5B6770'; ORANGE='F4B183'; GREEN='E2F0D9'
thin=Side(style='thin', color='D9E2E3')

wb=Workbook()
ws=wb.active; ws.title='00_LIRE_D_ABORD'

def title(ws, text, subtitle=None):
    ws.merge_cells('A1:H1'); ws['A1']=text; ws['A1'].font=Font(size=18,bold=True,color=WHITE); ws['A1'].fill=PatternFill('solid',fgColor=NAVY); ws['A1'].alignment=Alignment(vertical='center'); ws.row_dimensions[1].height=30
    if subtitle:
        ws.merge_cells('A2:H2'); ws['A2']=subtitle; ws['A2'].font=Font(size=10,italic=True,color=GREY); ws['A2'].fill=PatternFill('solid',fgColor=PALE); ws['A2'].alignment=Alignment(wrap_text=True,vertical='center'); ws.row_dimensions[2].height=30

def header(ws,row,cols):
    for col,name in enumerate(cols,1):
        c=ws.cell(row,col,name); c.font=Font(bold=True,color=WHITE); c.fill=PatternFill('solid',fgColor=TEAL); c.alignment=Alignment(wrap_text=True,vertical='center'); c.border=Border(bottom=thin)
    ws.row_dimensions[row].height=28

def fit(ws,widths):
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='A4'; ws.sheet_view.showGridLines=False
    for row in ws.iter_rows():
        for c in row:
            c.alignment=Alignment(vertical='top',wrap_text=True)

# 00

title(ws,'CLARTÉ360 — RÉFÉRENTIEL D’INTERPRÉTATION PIP RIASEC','Version 1.0 — source méthodologique du moteur d’interprétation à partir du Jalon E3')
rows=[
('Objet','Externaliser et documenter les règles d’interprétation du PIP. Le code ne doit plus contenir silencieusement les seuils et textes métier.'),
('Version référentiel','PIP-INT-1.0'),
('Statut','ACTIF — règles reprises à l’identique du moteur E2, désormais gouvernées par ce classeur.'),
('Principe','Une interprétation combine le niveau absolu de chaque score, le rang, les écarts entre dimensions et le relief global du profil.'),
('Important','Les seuils sont des repères descriptifs Clarté360 ancrés dans l’échelle PIP. Ils ne sont pas des normes de population, des seuils diagnostiques ou des niveaux d’aptitude.'),
('Source du score','Réponse 1–5 ; indice PIP = 25 × (moyenne − 1), donc 1→0, 2→25, 3→50, 4→75, 5→100.'),
('Gouvernance','Toute modification méthodologique doit être faite ici, versionnée, régénérer le JSON runtime et faire passer les tests. Ne pas modifier les règles directement dans interpretation.py.'),
('Portée','PIP Clarté360 uniquement. Les règles de ce référentiel ne s’appliquent pas aux scores O*NET.'),
]
header(ws,4,['Champ','Valeur'])
for r,(a,b) in enumerate(rows,5): ws.cell(r,1,a).font=Font(bold=True,color=NAVY); ws.cell(r,2,b)
fit(ws,[24,100,2,2,2,2,2,2])

# 01 dimensions
ws=wb.create_sheet('01_DIMENSIONS')
title(ws,'01 — Dimensions RIASEC','Libellés utilisés par le moteur et les rapports Clarté360')
header(ws,4,['code','libelle','description'])
dims={
'R':('Réaliste',"Attirance pour les activités concrètes, techniques et pratiques, où l’on agit sur des objets, des matières, des équipements, le vivant ou un environnement réel."),
'I':('Investigateur',"Attirance pour l’observation, l’analyse, la recherche d’informations, l’expérimentation, la compréhension et la résolution de problèmes."),
'A':('Artistique',"Attirance pour l’imagination, la création, l’expression et les activités laissant une place à l’originalité et aux choix personnels."),
'S':('Social',"Attirance pour les activités où l’on aide, transmet, conseille, accompagne, soigne ou rend un service directement utile à d’autres personnes."),
'E':('Entreprenant',"Attirance pour l’initiative, la persuasion, la négociation, la décision, la direction et la mobilisation autour d’objectifs."),
'C':('Conventionnel',"Attirance pour l’organisation, la fiabilité, le suivi, les règles, les données, les procédures et les processus structurés."),
}
for i,(code,(lab,desc)) in enumerate(dims.items(),5): ws.append([code,lab,desc])
fit(ws,[10,22,105])

# 02 bands
ws=wb.create_sheet('02_NIVEAUX_SCORES')
title(ws,'02 — Niveaux descriptifs des scores','Lecture du niveau absolu sur l’échelle PIP 0–100. Ce ne sont pas des normes psychométriques.')
header(ws,4,['rule_id','min_inclus','max_inclus','libelle','texte_autorise','statut'])
bands=[
('B01',0,24.999,'très peu marqué','vos réponses traduisent globalement peu d’attirance pour les activités de cette famille','ACTIF'),
('B02',25,44.999,'plutôt peu marqué','l’attirance exprimée reste plutôt limitée, même si certaines activités peuvent vous intéresser','ACTIF'),
('B03',45,59.999,'intermédiaire','vos réponses se situent autour d’une zone partagée, sans attirance ni rejet fortement affirmé','ACTIF'),
('B04',60,79.999,'marqué','vos réponses traduisent une attirance assez nette pour cette famille d’activités','ACTIF'),
('B05',80,100.001,'très marqué','vos réponses traduisent une attirance très nette pour cette famille d’activités','ACTIF'),
]
for row in bands: ws.append(row)
fit(ws,[12,14,14,24,95,12])

# 03 relief
ws=wb.create_sheet('03_RELIEF_PROFIL')
title(ws,'03 — Relief global du profil','Règles fondées sur l’amplitude entre le score le plus élevé et le plus faible.')
header(ws,4,['rule_id','spread_min_inclus','spread_max_inclus','libelle','texte_autorise','statut'])
reliefs=[
('R01',0,10,'très homogène','Les six dimensions sont très proches : le classement donne peu de relief au profil. Il est préférable de regarder les nuances, les facettes et votre expérience plutôt que de chercher une dominante unique.','ACTIF'),
('R02',10.000001,24,'peu contrasté','Les écarts entre dimensions restent limités : plusieurs familles d’activités peuvent coexister sans qu’une seule structure fortement le profil.','ACTIF'),
('R03',24.000001,44,'contrasté','Le profil présente des écarts suffisamment visibles pour distinguer des zones d’intérêt plus et moins présentes, tout en conservant plusieurs pôles possibles.','ACTIF'),
('R04',44.000001,100.001,'très contrasté','Le profil présente de forts écarts entre les dimensions les plus et les moins élevées. Certaines familles d’activités ressortent donc nettement davantage que d’autres.','ACTIF'),
]
for row in reliefs: ws.append(row)
fit(ws,[12,18,18,22,105,12])

# 04 head gaps
ws=wb.create_sheet('04_ECARTS_TETE')
title(ws,'04 — Écart entre les deux premières dimensions','Lecture de l’écart entre le 1er et le 2e score. Les textes utilisent des variables entre accolades.')
header(ws,4,['rule_id','gap_min_inclus','gap_max_inclus','libelle','texte_modele','statut'])
head=[
('T01',0,10,'double tête très proche','Les deux premières dimensions, {top_code} et {second_code}, sont très proches ({top_score:.0f} et {second_score:.0f}). Il est plus juste de les lire comme deux pôles associés que comme une première et une seconde nettement séparées.','ACTIF'),
('T02',10.000001,19,'première dimension légèrement détachée','La première dimension {top_code} ({top_score:.0f}) devance {second_code} ({second_score:.0f}), mais l’écart reste modéré. Les deux dimensions méritent d’être examinées ensemble.','ACTIF'),
('T03',19.000001,100.001,'première dimension nettement détachée','La première dimension {top_code} ({top_score:.0f}) se détache nettement de {second_code} ({second_score:.0f}). Elle donne davantage de relief au profil, sans constituer pour autant une aptitude ou une prescription professionnelle.','ACTIF'),
]
for row in head: ws.append(row)
fit(ws,[12,16,16,32,110,12])

# 05 scenarios
ws=wb.create_sheet('05_SCENARIOS_SYNTHESE')
title(ws,'05 — Scénarios de synthèse','Le premier scénario ACTIF dont toutes les conditions sont satisfaites est appliqué. Ordre = priorité croissante.')
header(ws,4,['priorite','rule_id','top_lt','top_gte','second_gte','gap12_lte','gap12_gte','headline_modele','summary_modele','statut'])
scenarios=[
(10,'S01',45,None,None,None,None,'{top_code} arrive en tête, mais aucune forte dominante ne ressort','{top_code} - {top_label} est votre dimension la plus élevée ({top_score:.0f}/100), mais ce niveau reste plutôt bas sur l’échelle PIP. Le premier rang signifie donc surtout qu’elle est la plus élevée relativement aux autres dimensions ; il ne signifie pas qu’il s’agit d’un intérêt fortement affirmé.','ACTIF'),
(20,'S02',None,60,60,10,None,'Deux dimensions ressortent à un niveau proche','{top_code} - {top_label} ({top_score:.0f}/100) et {second_code} - {second_label} ({second_score:.0f}/100) forment deux pôles d’intérêt proches. Il serait réducteur de présenter la première comme une dominante unique.','ACTIF'),
(30,'S03',None,80,None,None,15,'Une dominante {top_label_lower} très marquée se détache','{top_code} - {top_label} atteint {top_score:.0f}/100 et se détache nettement de la dimension suivante. Cela traduit une attirance déclarée particulièrement forte pour cette famille d’activités, et non une compétence ou une aptitude supérieure.','ACTIF'),
(99,'S99',None,None,None,None,None,'{top_code} - {top_label} constitue le premier pôle du profil','{top_code} - {top_label} arrive en tête à {top_score:.0f}/100. Son sens dépend à la fois de ce niveau absolu et de l’écart avec les autres dimensions : un rang ne s’interprète jamais seul.','ACTIF'),
]
for row in scenarios: ws.append(row)
fit(ws,[10,12,12,12,14,14,14,50,115,12])

# 06 safeguards
ws=wb.create_sheet('06_GARDE_FOUS')
title(ws,'06 — Garde-fous méthodologiques','Formulations à respecter dans toute évolution du moteur et du rapport.')
header(ws,4,['id','regle','consequence_pour_le_rapport','statut'])
guards=[
('G01','Le PIP mesure des intérêts professionnels, pas des compétences.','Ne jamais transformer un score élevé en aptitude, capacité ou niveau de compétence.','ACTIF'),
('G02','Le classement relatif ne remplace pas le niveau absolu.','Une première dimension à 31/100 ne doit pas être décrite comme une forte dominante.','ACTIF'),
('G03','Les seuils sont descriptifs Clarté360 et non normatifs.','Ne pas employer percentile, norme, population de référence, diagnostic ou seuil clinique.','ACTIF'),
('G04','Les écarts et quasi-égalités doivent être visibles.','Ne pas forcer une dominante unique lorsque les deux premières dimensions sont proches.','ACTIF'),
('G05','Les 30 facettes sont des indicateurs descriptifs.','Avec 2 ou 3 items par facette, ne pas les présenter comme 30 sous-échelles psychométriques validées.','ACTIF'),
('G06','Le PIP ne prescrit pas de métier.','Les rapprochements professionnels servent à explorer, jamais à recommander automatiquement un métier.','ACTIF'),
('G07','Ce référentiel ne s’applique pas à O*NET.','Ne pas appliquer les bandes 0–100 PIP aux scores O*NET.','ACTIF'),
]
for row in guards: ws.append(row)
fit(ws,[10,72,100,12])

# 07 examples
ws=wb.create_sheet('07_EXEMPLES_RECETTE')
title(ws,'07 — Exemples de recette','Cas lisibles humainement et exécutables par les tests. Les six scores sont saisis explicitement.')
header(ws,4,['case_id','R','I','A','S','E','C','ordre_attendu','scenario_attendu','relief_attendu','commentaire'])
examples=[
('EX01',31,28,26,24,22,20,'R,I,A,S,E,C','S01','peu contrasté','Première dimension basse : ne pas parler de forte dominante.'),
('EX02',35,48,87,93,42,30,'S,A,I,E,R,C','S02','très contrasté','Deux pôles élevés et proches : 93/87.'),
('EX03',30,41,52,93,35,20,'S,A,I,E,R,C','S03','très contrasté','Dominante élevée très détachée : 93/52.'),
('EX04',62,60,58,57,55,53,'R,I,A,S,E,C','S02','très homogène','Profil très homogène ; deux premières proches mais niveaux marqués.'),
('EX05',78,65,50,45,35,25,'R,I,A,S,E,C','S99','très contrasté','Premier pôle marqué, écart intermédiaire avec le second.'),
]
for row in examples: ws.append(row)
fit(ws,[12,8,8,8,8,8,8,28,18,22,80])

# 08 changelog
ws=wb.create_sheet('08_CHANGELOG')
title(ws,'08 — Changelog du référentiel','Toute modification publiée doit être tracée ici.')
header(ws,4,['date','version','type','description','justification','auteur_validateur'])
ws.append(['18/09/2026','PIP-INT-1.0','CREATION','Externalisation des règles du moteur E2 vers un référentiel méthodologique autonome et versionné.','Rendre les seuils, scénarios, textes et garde-fous lisibles, modifiables, auditables et testables sans les enfermer dans le code.','Clarté360'])
fit(ws,[16,18,18,105,100,28])

# common formatting
for ws in wb.worksheets:
    ws.auto_filter.ref = ws.dimensions if ws.max_row >= 4 and ws.max_column > 1 else None
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.border=Border(bottom=thin)
            if c.row % 2 == 1 and c.row>4: c.fill=PatternFill('solid',fgColor='FAFCFC')
    ws.freeze_panes='A5'

OUT.parent.mkdir(parents=True,exist_ok=True)
wb.save(OUT)
print(OUT)
