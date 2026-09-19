from __future__ import annotations
import argparse, hashlib, json
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'docs/sources/REFERENTIEL_INTERPRETATION_PIP_RIASEC_CLARTE360_V1_0.xlsx'
OUT=ROOT/'resources/runtime/pip_interpretation_PIP-INT-1.0.json'


def sha256(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''): h.update(chunk)
    return h.hexdigest()

def recs(ws):
    raw=[c.value for c in ws[4]]
    indexes=[i for i,h in enumerate(raw) if h is not None]
    headers=[raw[i] for i in indexes]
    out=[]
    for vals in ws.iter_rows(min_row=5,values_only=True):
        if not any(v is not None for v in vals): continue
        out.append({headers[j]:vals[indexes[j]] for j in range(len(headers))})
    return out

def active(rows): return [r for r in rows if str(r.get('statut') or '').upper()=='ACTIF']

def build(source=SRC):
    wb=openpyxl.load_workbook(source,data_only=True,read_only=True)
    required={'00_LIRE_D_ABORD','01_DIMENSIONS','02_NIVEAUX_SCORES','03_RELIEF_PROFIL','04_ECARTS_TETE','05_SCENARIOS_SYNTHESE','06_GARDE_FOUS','07_EXEMPLES_RECETTE','08_CHANGELOG'}
    if not required.issubset(wb.sheetnames): raise ValueError(f'Onglets manquants: {sorted(required-set(wb.sheetnames))}')
    dims=recs(wb['01_DIMENSIONS'])
    if [r['code'] for r in dims] != list('RIASEC'): raise ValueError('Ordre RIASEC invalide')
    labels={r['code']:r['libelle'] for r in dims}; descriptions={r['code']:r['description'] for r in dims}
    bands=active(recs(wb['02_NIVEAUX_SCORES'])); relief=active(recs(wb['03_RELIEF_PROFIL'])); head=active(recs(wb['04_ECARTS_TETE'])); scenarios=active(recs(wb['05_SCENARIOS_SYNTHESE'])); guards=active(recs(wb['06_GARDE_FOUS'])); examples=recs(wb['07_EXEMPLES_RECETTE'])
    if len(bands)!=5 or len(relief)!=4 or len(head)!=3 or len(scenarios)<1: raise ValueError('Nombre de règles inattendu')
    # Coverage checks
    if float(bands[0]['min_inclus'])!=0 or float(bands[-1]['max_inclus'])<100: raise ValueError('Couverture scores 0-100 incomplète')
    scenarios=sorted(scenarios,key=lambda r:int(r['priorite']))
    if str(scenarios[-1]['rule_id'])!='S99': raise ValueError('Le scénario de repli S99 doit être le dernier')
    return {
        'schema_version':'PIP-INTERPRETATION-RUNTIME-1.0',
        'interpretation_version':'PIP-INT-1.0',
        'source_file':source.name,
        'source_sha256':sha256(source),
        'dimension_order':list('RIASEC'),
        'labels':labels,
        'descriptions':descriptions,
        'bands':bands,
        'relief_rules':relief,
        'head_gap_rules':head,
        'summary_scenarios':scenarios,
        'safeguards':guards,
        'examples':examples,
    }

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--check',action='store_true'); a=ap.parse_args(); data=build()
    if a.check:
        if not OUT.exists(): raise SystemExit(f'Runtime absent: {OUT}')
        cur=json.loads(OUT.read_text(encoding='utf-8'))
        if cur!=data: raise SystemExit('ÉCART RÉFÉRENTIEL/RUNTIME: régénérer le JSON d’interprétation')
        print('OK — référentiel d’interprétation et runtime JSON sont synchronisés.')
        return 0
    OUT.parent.mkdir(parents=True,exist_ok=True)
    OUT.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'Généré: {OUT.relative_to(ROOT)}')
    return 0
if __name__=='__main__': raise SystemExit(main())
