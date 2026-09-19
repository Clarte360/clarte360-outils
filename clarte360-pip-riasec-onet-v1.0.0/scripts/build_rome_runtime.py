from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/references/rome_riasec_clarte360.xlsx'
RUNTIME = ROOT / 'resources/runtime/rome_riasec_ROME-RIASEC-2026-06.json'
VERSION = 'ROME-RIASEC-2026-06'


def _clean(v):
    return '' if v is None else str(v).strip()


def build_payload() -> dict:
    wb = load_workbook(SOURCE, read_only=False, data_only=True)
    ws = wb['ROME_RIASEC']
    rows = list(ws.iter_rows(values_only=True))
    headers = [_clean(v) for v in rows[0]]
    idx = {h: i for i, h in enumerate(headers)}
    required = ['Code ROME', 'Intitule fiche', 'RIASEC normalise', 'PDF source', 'Statut extraction']
    missing = [h for h in required if h not in idx]
    if missing:
        raise ValueError(f'Colonnes ROME manquantes: {missing}')

    records = []
    for row in rows[1:]:
        code = _clean(row[idx['Code ROME']])
        title = _clean(row[idx['Intitule fiche']])
        profile = _clean(row[idx['RIASEC normalise']]).upper()
        status = _clean(row[idx['Statut extraction']]).upper()
        if not code or not title or status != 'OK':
            continue
        records.append({
            'rome_code': code,
            'title': title,
            'riasec_profile': profile,
            'source_pdf': _clean(row[idx['PDF source']]),
        })

    # Preserve source order for traceability. Selection logic can sort deterministically later.
    controls = wb['CONTROLES']
    date_rome = None
    source_name = None
    for r in controls.iter_rows(values_only=True):
        if _clean(r[0]) == 'Date ROME':
            date_rome = _clean(r[1])
        if _clean(r[0]) == 'Source':
            source_name = _clean(r[1])

    return {
        'reference_version': VERSION,
        'source_file': SOURCE.name,
        'source_date': date_rome or 'juin 2026',
        'source_archive': source_name or '',
        'record_count': len(records),
        'records': records,
        'usage_rule': 'Exploration uniquement. Ne jamais présenter ces métiers comme des recommandations ou prescriptions.',
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    args = ap.parse_args()
    payload = build_payload()
    text = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=False) + '\n'
    if args.check:
        if not RUNTIME.exists():
            print('Runtime ROME absent:', RUNTIME)
            return 1
        current = RUNTIME.read_text(encoding='utf-8')
        if current != text:
            print('Runtime ROME non synchronisé avec le tableur source.')
            return 1
        print(f'ROME OK: {payload["record_count"]} fiches, version {payload["reference_version"]}')
        return 0
    RUNTIME.parent.mkdir(parents=True, exist_ok=True)
    RUNTIME.write_text(text, encoding='utf-8')
    print(f'Écrit: {RUNTIME} ({payload["record_count"]} fiches)')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
