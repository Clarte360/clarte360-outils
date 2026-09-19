from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs/sources/CLARTE360_PIP_RIASEC_TABLEUR_MAITRE_V0_6_ACTIF.xlsx"
OUT = ROOT / "resources/runtime/pip_bank_PIP-BANK-0.5.json"
DIMS = tuple("RIASEC")
FACETS = {f"{d}{n}" for d in DIMS for n in range(1, 6)}


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _sheet_records(ws) -> list[dict]:
    headers = [c.value for c in ws[1]]
    records = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None for v in values):
            continue
        records.append({headers[i]: values[i] for i in range(len(headers))})
    return records


def build_runtime(source: Path = SRC) -> dict:
    if not source.exists():
        raise FileNotFoundError(source)
    wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    if "01_QUESTIONS_72" not in wb.sheetnames or "02_REPONSES_CONSIGNE" not in wb.sheetnames:
        raise ValueError("Le tableur maître doit contenir 01_QUESTIONS_72 et 02_REPONSES_CONSIGNE.")

    qrows = _sheet_records(wb["01_QUESTIONS_72"])
    active = [r for r in qrows if str(r.get("statut") or "").upper() == "ACTIF"]
    if len(active) != 72:
        raise ValueError(f"72 items ACTIF attendus, reçu {len(active)}")

    ids = [str(r["item_id"]).strip() for r in active]
    if len(ids) != len(set(ids)):
        raise ValueError("IDs d'items dupliqués dans le tableur maître.")

    dims = Counter(str(r["dimension"]).strip() for r in active)
    if dims != Counter({d: 12 for d in DIMS}):
        raise ValueError(f"Répartition RIASEC invalide: {dict(dims)}")

    facets = Counter(str(r["facette"]).strip() for r in active)
    if set(facets) != FACETS or min(facets.values()) < 2:
        raise ValueError("Les 30 facettes doivent être couvertes par au moins 2 items.")

    versions = {str(r["item_version"]).strip() for r in active}
    banks = {str(r["bank_version"]).strip() for r in active}
    if versions != {"0.6"}:
        raise ValueError(f"item_version attendue 0.6, reçu {versions}")
    if banks != {"PIP-BANK-0.5"}:
        raise ValueError(f"bank_version attendue PIP-BANK-0.5, reçu {banks}")

    protocol = wb["02_REPONSES_CONSIGNE"]
    instruction = str(protocol["B4"].value or "").strip()
    if not instruction:
        raise ValueError("Consigne de réponse absente en 02_REPONSES_CONSIGNE!B4")

    scale = {}
    header_row = None
    for row in range(1, protocol.max_row + 1):
        if protocol.cell(row, 1).value == "score" and protocol.cell(row, 2).value == "libellé affiché":
            header_row = row
            break
    if header_row is None:
        raise ValueError("Table de l'échelle de réponse introuvable.")
    for row in range(header_row + 1, header_row + 6):
        score = protocol.cell(row, 1).value
        label = protocol.cell(row, 2).value
        if score is None or not label:
            raise ValueError("Échelle de réponse incomplète.")
        scale[str(int(score))] = str(label).strip()
    if set(scale) != {"1", "2", "3", "4", "5"}:
        raise ValueError("Échelle de réponse 1–5 invalide.")

    items = []
    for r in sorted(active, key=lambda x: int(x["ordre"])):
        text = str(r["question_finale"] or "").strip()
        if len(text) < 20:
            raise ValueError(f"Question trop courte ou vide: {r['item_id']}")
        items.append({
            "item_id": str(r["item_id"]).strip(),
            "item_version": str(r["item_version"]).strip(),
            "statut": "ACTIF",
            "bloc": str(r["bloc"]).strip(),
            "dimension": str(r["dimension"]).strip(),
            "facette": str(r["facette"]).strip(),
            "texte_fr": text,
        })

    return {
        "schema_version": "PIP-RUNTIME-1.1",
        "bank_version": "PIP-BANK-0.5",
        "source_file": source.name,
        "source_sha256": _sha256(source),
        "dimensions_order": list(DIMS),
        "blocks_order": ["ACT", "SIT", "ENV"],
        "response_instruction": instruction,
        "response_scale": scale,
        "items": items,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Génère la banque runtime PIP depuis le tableur maître actif.")
    parser.add_argument("--check", action="store_true", help="Vérifie que le JSON versionné correspond exactement au tableur maître.")
    args = parser.parse_args()

    runtime = build_runtime()
    if args.check:
        if not OUT.exists():
            raise SystemExit(f"Runtime absent: {OUT}")
        current = json.loads(OUT.read_text(encoding="utf-8"))
        if current != runtime:
            raise SystemExit("ÉCART SOURCE/RUNTIME: régénérer resources/runtime/pip_bank_PIP-BANK-0.5.json")
        print("OK — tableur maître et runtime JSON sont synchronisés.")
        return 0

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(runtime, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Généré: {OUT.relative_to(ROOT)}")
    print(f"Source: {SRC.relative_to(ROOT)}")
    print(f"Items: {len(runtime['items'])} | Banque: {runtime['bank_version']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
