from io import BytesIO

from docx import Document

from clarte360_cp.reporting.conclusion_docx import (
    render_action_plan_docx,
    render_results_docx,
    render_synthesis_docx,
)


def test_final_conclusion_docx_exports_are_editable_and_use_clarte_letterhead():
    outputs = (
        render_action_plan_docx({"version_no": 2, "items": [{
            "action": "Rencontrer deux professionnels du métier",
            "due": "Avant le prochain rendez-vous",
            "means": "Préparer les questions",
            "expected": "Deux retours documentés",
        }]}),
        render_results_docx({"version": 3, "content": {
            "parcours": "Expérience et études",
            "savoirs": "Comptabilité générale",
            "projet": "Chef comptable",
            "diploma_status": "Aucun diplôme obligatoire établi",
            "diploma_source": "Vérification officielle datée",
        }}),
        render_synthesis_docx({"version": 1, "content": {
            "objectif": "Bilan de compétences",
            "projet": "Chef comptable",
            "conditions": "Consolider les savoirs",
            "plan": "Enquête et mise en pratique",
        }}),
    )
    expected_titles = (
        "Plan d’action convenu",
        "Résultats détaillés du bilan",
        "Synthèse finale du bilan",
    )

    for output, expected_title in zip(outputs, expected_titles):
        document = Document(BytesIO(output))
        body = "\n".join(paragraph.text for paragraph in document.paragraphs)
        footer = " ".join(paragraph.text for paragraph in document.sections[0].footer.paragraphs)
        header = document.sections[0].header.paragraphs[0]
        assert expected_title in body
        assert "Document de travail confidentiel" in body
        assert "CLARTÉ360" in header.text
        assert "SIRET : 10234983400014" in footer
        assert abs(document.sections[0].page_width.cm - 21) < 0.01
        assert abs(document.sections[0].page_height.cm - 29.7) < 0.01
