"""Editable Word versions of finalized Clarté360 conclusion documents."""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from clarte360_cp.reporting.conclusion_content import (
    action_plan_sections, results_sections, synthesis_sections,
)

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "assets" / "logo_clarte360.png"


def render_conclusion_docx(title: str, sections: list[tuple[str, str]], *, version: str | int) -> bytes:
    """Generate an editable DOCX from an already confirmed/published artifact."""
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.7)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(36, 58, 58)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(8)
    title_properties = title_style._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 2"].font.size = Pt(12)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(14)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if LOGO.is_file():
        header.add_run().add_picture(str(LOGO), width=Inches(0.42))
        header.add_run("  ")
    brand = header.add_run("CLARTÉ360  |  Bilan de compétences")
    brand.bold = True
    brand.font.name = "Arial"
    brand.font.size = Pt(9)
    brand.font.color.rgb = RGBColor(8, 126, 128)

    footer = section.footer
    footer.paragraphs[0].text = (
        "CLARTÉ360 – 60 rue François 1er – 75008 Paris – Tél. : 01 89 48 08 25 – "
        "contact@Clarté360.com – www.Clarté360.com"
    )
    legal = footer.add_paragraph("RCS : 102349834 · SIRET : 10234983400014 · NAF : 8559 A · Id CEE : FR88102349834")
    for paragraph in footer.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(92, 112, 112)

    document.core_properties.title = str(title)[:160]
    document.core_properties.subject = "Bilan de compétences"
    document.core_properties.author = "Clarté360"
    document.core_properties.last_modified_by = "Clarté360"

    heading = document.add_paragraph(style="Title")
    heading.add_run(str(title))
    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(6)
    run = meta.add_run(f"Version {version} · Préparé le {date.today().isoformat()}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(92, 112, 112)
    confidentiality = document.add_paragraph("Document de travail confidentiel")
    confidentiality.paragraph_format.space_after = Pt(12)
    for run in confidentiality.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(92, 112, 112)

    for label, content in sections:
        document.add_paragraph(str(label), style="Heading 2")
        lines = str(content or "Non renseigné").splitlines() or ["Non renseigné"]
        for line in lines:
            document.add_paragraph(line or " ")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_action_plan_docx(plan: dict) -> bytes:
    return render_conclusion_docx("Plan d’action convenu", action_plan_sections(plan), version=plan.get("version_no") or 1)


def render_results_docx(record: dict) -> bytes:
    return render_conclusion_docx("Résultats détaillés du bilan", results_sections(record), version=record.get("version") or 1)


def render_synthesis_docx(record: dict) -> bytes:
    return render_conclusion_docx("Synthèse finale du bilan", synthesis_sections(record), version=record.get("version") or 1)
